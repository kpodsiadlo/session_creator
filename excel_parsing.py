import string
from openpyxl import load_workbook


def get_filenames_from_excel_column(
        filename, column_letter, row_start_num, row_stop_num):
    """Return a list of entries in a given column"""

    file_list = []
    if len(column_letter) != 1:
        raise ValueError("The column ID must be \
                        a single letter between A and Z")
    else:
        try:
            column_letter_lowercase = column_letter.lower()
        except:
            raise ValueError("The column ID must be \
                            a single letter between A and Z")
        finally:
            column_number = int(ord(column_letter_lowercase) - 96)

    # get the file
    wb = load_workbook(filename)
    # get the sheet
    ws = wb['Sheet1']
    # get the values
    file_list = [ws.cell(row=i, column=column_number).value
                 for i in range(row_start_num, row_stop_num)]

    return file_list
